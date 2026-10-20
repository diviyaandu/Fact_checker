"""Claim-level evaluation of the fact-checker's verdicts.

    python -m evaluation.run_eval path/to/claims.xlsx            # full run
    python -m evaluation.run_eval path/to/claims.xlsx --validate-only
    python -m evaluation.run_eval path/to/claims.xlsx --limit 5  # smoke test
    python -m evaluation.run_eval path/to/claims.xlsx --fresh    # ignore saved predictions

Each claim goes through the app's own evidence + verdict stage
(factcheck.pipeline.get_claim_verdicts); transcript and claim extraction are
skipped (no extra Groq call). Only Ground_Truth is used as the label; the
input file is never modified and the system's predictions never touch it.
"""

import argparse
import csv
import json
import os
import subprocess
import sys
import time
from datetime import datetime, timezone

from evaluation.metrics import LABELS, compute_metrics

DEFAULT_OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results")
CHUNK_SIZE = 10  # claims per save point (also bounds work lost on a crash)
COLUMNS = [
    "Claim_ID", "Claim", "Ground_Truth", "System_Prediction", "Confidence",
    "Correct", "Notes", "Sources_Found", "Queries_Succeeded", "Error",
]


# ---- dataset ---------------------------------------------------------------

def _read_rows(path: str) -> list[dict]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Reading .xlsx needs openpyxl:  pip install openpyxl  (or pass a .csv)")
        sheet = load_workbook(path, read_only=True, data_only=True).active
        grid = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            grid = list(csv.reader(f))

    grid = [r for r in grid if any(c not in (None, "") for c in r)]
    if not grid:
        sys.exit("Dataset is empty.")

    header = [str(h or "").strip().lower() for h in grid[0]]
    return [
        {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row)}
        for row in grid[1:]
    ]


def load_dataset(path: str) -> list[dict]:
    """Validated rows: {Claim_ID, Claim, Ground_Truth, Notes}. Exits on bad data."""
    rows = _read_rows(path)
    out, problems, seen = [], [], set()

    for n, r in enumerate(rows, start=2):  # 2 = first data row in the sheet
        claim_id, claim = r.get("claim_id", ""), r.get("claim", "")
        truth = r.get("ground_truth", "").upper()

        if not claim_id or not claim:
            problems.append(f"row {n}: missing Claim_ID or Claim")
        elif truth not in LABELS:
            problems.append(f"row {n}: Ground_Truth {r.get('ground_truth', '')!r} not in {LABELS}")
        elif claim_id in seen:
            problems.append(f"row {n}: duplicate Claim_ID {claim_id}")
        else:
            seen.add(claim_id)
            out.append({"Claim_ID": claim_id, "Claim": claim,
                        "Ground_Truth": truth, "Notes": r.get("notes", "")})

    if problems:
        sys.exit("Dataset problems:\n  " + "\n  ".join(problems))
    return out


# ---- saved predictions (also the resume store) -----------------------------

def _write_csv(path: str, header: list[str], rows: list[list]) -> None:
    tmp = path + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    os.replace(tmp, path)


def save_predictions(path: str, rows: list[dict]) -> None:
    _write_csv(path, COLUMNS, [[r.get(c, "") for c in COLUMNS] for r in rows])


def load_saved(path: str) -> dict[str, dict]:
    if not os.path.exists(path):
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {r["Claim_ID"]: r for r in csv.DictReader(f)}


def _reusable(saved: dict | None, item: dict) -> bool:
    """A saved row is reused only if it is for the same claim/label and had no error."""
    return bool(
        saved and saved.get("System_Prediction") and not saved.get("Error")
        and saved.get("Claim") == item["Claim"]
        and saved.get("Ground_Truth") == item["Ground_Truth"]
    )


# ---- running the existing pipeline -----------------------------------------

def _run_pipeline(items: list[dict], out_dir: str) -> dict[str, dict]:
    """Runs the claims through the app's evidence+verdict stage.
    Returns {Claim_ID: prediction fields}."""
    # Lazy imports: these need API keys, so --validate-only never loads them.
    from factcheck import cache
    from factcheck.models import Claim
    from factcheck.pipeline.verdicts import get_claim_verdicts
    from factcheck.verdicts.parsing import fallback_invalid_verdict

    # The pipeline always writes finished verdicts to its disk cache; point that
    # at an eval-only file so the app's cache is left untouched.
    cache._CACHE_PATH = os.path.join(out_dir, ".eval_cache.json")
    invalid_summary = fallback_invalid_verdict().summary

    claims = [
        Claim(
            claim_id=i, claim_text=it["Claim"], start_time="00:00", end_time="00:00",
            category="General", importance="medium",
            search_queries=[],  # empty -> the pipeline's built-in fallback queries
        )
        for i, it in enumerate(items, start=1)
    ]

    results = get_claim_verdicts(claims, use_cache=False)

    preds = {}
    for i, it in enumerate(items, start=1):
        res = results[i]
        verdict, meta = res["verdict"], res["search_meta"] or {}
        error = ""
        if res["search_error"]:
            error = f"search_error: {res['search_error']}"
        elif meta.get("queries_succeeded", 0) == 0:
            error = "search_failed"
        elif verdict.summary == invalid_summary:
            error = "llm_invalid_response"

        preds[it["Claim_ID"]] = {
            "System_Prediction": verdict.verdict,
            "Confidence": round(verdict.confidence, 4),
            "Sources_Found": meta.get("sources_found", 0),
            "Queries_Succeeded": meta.get("queries_succeeded", 0),
            "Error": error,
        }
    return preds


# ---- reporting --------------------------------------------------------------

def _git_commit() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"], stderr=subprocess.DEVNULL, text=True
        ).strip()
    except Exception:
        return "unknown"


def write_reports(out_dir: str, rows: list[dict]) -> dict:
    m = compute_metrics([r["Ground_Truth"] for r in rows], [r["System_Prediction"] for r in rows])
    n_err = sum(1 for r in rows if r.get("Error"))

    from factcheck import config
    from factcheck.scoring.constants import FORMULA_VERSION
    m["meta"] = {
        "timestamp_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "git_commit": _git_commit(),
        "verdict_model": config.VERDICT_MODEL,
        "batch_size": config.BATCH_SIZE,
        "confidence_formula_version": FORMULA_VERSION,
        "search_queries": "pipeline fallback queries (not LLM-generated)",
        "rows_with_errors": n_err,
    }
    with open(os.path.join(out_dir, "metrics.json"), "w", encoding="utf-8") as f:
        json.dump(m, f, indent=2)

    _write_csv(
        os.path.join(out_dir, "confusion_matrix.csv"),
        ["Truth \\ Predicted"] + LABELS,
        [[LABELS[i]] + row for i, row in enumerate(m["confusion_matrix"])],
    )
    return m


def print_report(m: dict) -> None:
    print(f"\n{'Class':<16}{'P':>7}{'R':>7}{'F1':>7}{'N':>5}")
    for label, c in m["per_class"].items():
        print(f"{label:<16}{c['precision']:>7.3f}{c['recall']:>7.3f}{c['f1']:>7.3f}{c['support']:>5}")
    print(f"\nMACRO F1: {m['macro_f1']:.4f}   accuracy: {m['accuracy']:.4f}   n={m['n']}")
    print("\nConfusion matrix (rows=truth, cols=predicted):")
    w = max(len(l) for l in m["labels"]) + 2
    print(f"{'':<16}" + "".join(f"{l:>{w}}" for l in m["labels"]))
    for label, row in zip(m["labels"], m["confusion_matrix"]):
        print(f"{label:<16}" + "".join(f"{v:>{w}}" for v in row))


# ---- main -------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("dataset", help=".xlsx or .csv with Claim_ID, Claim, Ground_Truth")
    ap.add_argument("--out", default=DEFAULT_OUT, help="results directory")
    ap.add_argument("--limit", type=int, help="only the first N claims (smoke test; metrics are partial)")
    ap.add_argument("--fresh", action="store_true", help="ignore saved predictions and re-run everything")
    ap.add_argument("--validate-only", action="store_true", help="check the dataset; make no API calls")
    ap.add_argument("--pause", type=float, default=0, help="seconds to wait between chunks (rate limits)")
    args = ap.parse_args()

    items = load_dataset(args.dataset)
    if args.limit:
        items = items[: args.limit]

    counts = {l: sum(1 for it in items if it["Ground_Truth"] == l) for l in LABELS}
    print(f"{len(items)} claims: {counts}")
    if args.validate_only:
        return

    os.makedirs(args.out, exist_ok=True)
    pred_path = os.path.join(args.out, "predictions.csv")
    saved = {} if args.fresh else load_saved(pred_path)

    rows = {it["Claim_ID"]: {**it} for it in items}
    todo = []
    for it in items:
        if _reusable(saved.get(it["Claim_ID"]), it):
            rows[it["Claim_ID"]].update(saved[it["Claim_ID"]])
        else:
            todo.append(it)
    print(f"reusing {len(items) - len(todo)} saved predictions; running {len(todo)}")

    def finish(r: dict) -> dict:
        r["Correct"] = "TRUE" if r.get("System_Prediction") == r["Ground_Truth"] else "FALSE"
        return r

    for start in range(0, len(todo), CHUNK_SIZE):
        chunk = todo[start: start + CHUNK_SIZE]
        for cid, pred in _run_pipeline(chunk, args.out).items():
            rows[cid].update(pred)
            finish(rows[cid])
        done = [finish(rows[it["Claim_ID"]]) for it in items if rows[it["Claim_ID"]].get("System_Prediction")]
        save_predictions(pred_path, done)
        print(f"  {len(done)}/{len(items)} done")
        if args.pause and start + CHUNK_SIZE < len(todo):
            time.sleep(args.pause)

    final = [finish(rows[it["Claim_ID"]]) for it in items]
    save_predictions(pred_path, final)
    m = write_reports(args.out, final)
    print_report(m)

    errs = m["meta"]["rows_with_errors"]
    if errs:
        print(f"\nWARNING: {errs} claim(s) hit a search/LLM error (see Error column); they are scored as "
              "predicted. Re-run the same command to retry only those.")
    print(f"\nSaved to {args.out}: predictions.csv, metrics.json, confusion_matrix.csv")


if __name__ == "__main__":
    main()